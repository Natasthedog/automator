from __future__ import annotations

# Auto-generated split from dash_app.py

import io
import base64
import logging
import copy
from difflib import SequenceMatcher
from dataclasses import dataclass, asdict, is_dataclass
from datetime import date, timedelta
from pathlib import Path
import re
import json
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries, get_column_letter
import numbers
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.parts.embeddedpackage import EmbeddedXlsxPart
from lxml import etree

logger = logging.getLogger(__name__)

from ..pptx.slides import _slide_title
from .targets import _normalize_column_name

def _normalize_text_value(value) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip().lower()


def _normalize_lookup_value(value: str) -> str:
    normalized = str(value).strip().casefold()
    normalized = re.sub(r"[^\w\s]", " ", normalized)
    normalized = normalized.replace("_", " ")
    normalized = " ".join(normalized.split())
    return normalized


def _format_fuzzy_candidates(candidates: list[tuple[float, str]]) -> str:
    return ", ".join(f"{label!r} ({score:.1f})" for score, label in candidates)


def _resolve_column_from_candidates(
    df: pd.DataFrame,
    header_row: pd.Series | None,
    candidates: list[str],
    *,
    context: str,
    threshold: float = 85.0,
) -> tuple[str | None, int, float]:
    column_options: list[tuple[str, object, int]] = []
    for column in df.columns:
        column_options.append((str(column), column, 0))
    if header_row is not None:
        for column, value in header_row.items():
            if pd.isna(value):
                continue
            column_options.append((str(value), column, 1))

    normalized_candidates = [
        (candidate, _normalize_lookup_value(candidate)) for candidate in candidates
    ]
    normalized_options = [
        (label, _normalize_lookup_value(label), column, data_start_idx)
        for label, column, data_start_idx in column_options
        if _normalize_lookup_value(label)
    ]

    exact_matches = []
    for candidate, normalized_candidate in normalized_candidates:
        for label, normalized_label, column, data_start_idx in normalized_options:
            if normalized_candidate == normalized_label:
                exact_matches.append((candidate, label, column, data_start_idx, 100.0))
    if exact_matches:
        unique_columns = {match[2] for match in exact_matches}
        if len(unique_columns) > 1:
            top_candidates = [
                (match[4], match[1]) for match in exact_matches[:5]
            ]
            raise ValueError(
                f"Ambiguous {context} match. Top candidates: "
                f"{_format_fuzzy_candidates(top_candidates)}"
            )
        candidate, label, column, data_start_idx, score = exact_matches[0]
        logger.info(
            'Resolved header "%s" -> "%s" (score %.1f)',
            candidate,
            label,
            score,
        )
        return column, data_start_idx, score

    from difflib import SequenceMatcher

    scored: list[tuple[float, str, object, int, str]] = []
    for candidate, normalized_candidate in normalized_candidates:
        for label, normalized_label, column, data_start_idx in normalized_options:
            score = SequenceMatcher(None, normalized_candidate, normalized_label).ratio() * 100
            scored.append((score, label, column, data_start_idx, candidate))
    if not scored:
        return None, 0, 0.0
    scored.sort(key=lambda item: item[0], reverse=True)
    top_score = scored[0][0]
    if top_score < threshold:
        return None, 0, top_score
    close_matches = [
        (score, label, column)
        for score, label, column, _, _ in scored
        if score >= top_score - 1.0
    ]
    unique_columns = {column for _, _, column in close_matches}
    if len(unique_columns) > 1:
        top_candidates = [(score, label) for score, label, _ in close_matches[:5]]
        raise ValueError(
            f"Ambiguous {context} match. Top candidates: "
            f"{_format_fuzzy_candidates(top_candidates)}"
        )
    score, label, column, data_start_idx, candidate = scored[0]
    logger.info(
        'Resolved header "%s" -> "%s" (score %.1f)',
        candidate,
        label,
        score,
    )
    return column, data_start_idx, score


def _resolve_label_from_text(text: str, labels: list[str], threshold: float = 85.0) -> str:
    normalized_text = _normalize_lookup_value(text)
    if not normalized_text:
        raise ValueError("Slide text is empty after normalization.")
    normalized_labels = []
    exact_matches = []
    for label in labels:
        normalized_label = _normalize_lookup_value(label)
        normalized_labels.append((label, normalized_label))
        if normalized_label == normalized_text:
            exact_matches.append(label)
    if exact_matches:
        resolved = exact_matches[0]
        logger.info(
            "Resolved slide text %r -> %r (score 100.0)",
            text,
            resolved,
        )
        return resolved
    from difflib import SequenceMatcher

    scored = []
    for label, normalized_label in normalized_labels:
        if not normalized_label:
            continue
        score = SequenceMatcher(None, normalized_text, normalized_label).ratio() * 100
        scored.append((score, label))
    scored.sort(reverse=True)
    top_candidates = scored[:5]
    if not scored or scored[0][0] < threshold:
        raise ValueError(
            "No slide text match found (threshold {:.0f}). Top candidates: {}".format(
                threshold,
                _format_fuzzy_candidates(top_candidates),
            )
        )
    best_score, best_label = scored[0]
    if len(scored) > 1 and scored[1][0] >= threshold and (best_score - scored[1][0]) < 2:
        raise ValueError(
            "Ambiguous slide text match (threshold {:.0f}). Top candidates: {}".format(
                threshold,
                _format_fuzzy_candidates(top_candidates),
            )
        )
    logger.info(
        "Resolved slide text %r -> %r (score %.1f)",
        text,
        best_label,
        best_score,
    )
    return best_label


def _resolve_target_level_label_for_slide(slide, labels: list[str]) -> str | None:
    candidates = []
    slide_title = _slide_title(slide)
    if slide_title:
        candidates.append(("title", slide_title))
    slide_name = getattr(slide, "name", None) or ""
    if slide_name:
        candidates.append(("name", slide_name))
    if not candidates:
        return None
    errors = []
    for source, text in candidates:
        try:
            return _resolve_label_from_text(text, labels)
        except ValueError as exc:
            message = str(exc)
            error_message = (
                f"Could not resolve Target Level Label from slide {source} {text!r}: {exc}"
            )
            if message.startswith("No slide text match found") or message.startswith(
                "Slide text is empty after normalization"
            ):
                logger.info("%s", error_message)
                continue
            errors.append(error_message)
    if errors:
        raise ValueError(" | ".join(errors))
    return None


def _resolve_target_label_for_slide(slide, labels: list[str]) -> str | None:
    candidates = []
    slide_title = _slide_title(slide)
    if slide_title:
        candidates.append(("title", slide_title))
    slide_name = getattr(slide, "name", None) or ""
    if slide_name:
        candidates.append(("name", slide_name))
    if not candidates:
        return None
    errors = []
    for source, text in candidates:
        try:
            return _resolve_label_from_text(text, labels)
        except ValueError as exc:
            error_message = f"Could not resolve Target Label from slide {source} {text!r}: {exc}"
            errors.append(error_message)
    if errors:
        raise ValueError(" | ".join(errors))
    return None


def _normalize_target_label_value(value: str | None) -> str:
    normalized = _normalize_text_value(value)
    if normalized == "competitor":
        return "cross"
    return normalized


def _two_row_column_match(
    group_value: str,
    sub_value: str,
    candidates: list[str],
) -> bool:
    group_key = _normalize_column_name(group_value)
    sub_key = _normalize_column_name(sub_value)
    candidate_keys = {_normalize_column_name(candidate) for candidate in candidates}
    return group_key in candidate_keys or sub_key in candidate_keys


def _parse_two_row_header_dataframe(
    raw_df: pd.DataFrame,
) -> tuple[pd.DataFrame, dict]:
    """Parse a gatheredCN10 file that uses two header rows.

    Returns the data rows with stable internal column IDs plus metadata for UI mapping.

    Example:
        >>> raw = pd.DataFrame(
        ...     [
        ...         ["Promo", "Promo", "", ""],
        ...         ["Feature", "Display", "Target Label", "Year"],
        ...         [1, 2, "Own", 2023],
        ...     ]
        ... )
        >>> data_df, meta = _parse_two_row_header_dataframe(raw)
        >>> meta["group_order"]
        ['Promo']
    """
    if raw_df is None or raw_df.empty or raw_df.shape[0] < 3:
        raise ValueError("The gatheredCN10 file must include two header rows and data rows.")
    header_row1 = raw_df.iloc[0].fillna("")
    header_row2 = raw_df.iloc[1].fillna("")
    columns_meta = []
    group_map: dict[str, list[dict]] = {}
    group_order: list[str] = []
    for idx in range(raw_df.shape[1]):
        group = str(header_row1.iloc[idx]).strip()
        subheader = str(header_row2.iloc[idx]).strip()
        col_id = f"col_{idx}"
        columns_meta.append(
            {
                "id": col_id,
                "group": group,
                "subheader": subheader,
                "position": idx,
            }
        )
        if not group:
            continue
        group_key = _normalize_column_name(group)
        if group_key in {"targetlabel", "year"}:
            continue
        if group not in group_map:
            group_map[group] = []
            group_order.append(group)
        group_map[group].append(
            {
                "id": col_id,
                "subheader": subheader,
                "position": idx,
            }
        )

    target_label_id = None
    year_id = None
    for column in columns_meta:
        if target_label_id is None and _two_row_column_match(
            column["group"],
            column["subheader"],
            ["Target Label"],
        ):
            target_label_id = column["id"]
        if year_id is None and _two_row_column_match(
            column["group"],
            column["subheader"],
            ["Year"],
        ):
            year_id = column["id"]

    data_df = raw_df.iloc[2:].reset_index(drop=True).copy()
    data_df.columns = [col["id"] for col in columns_meta]
    metadata = {
        "columns": columns_meta,
        "groups": group_map,
        "group_order": group_order,
        "target_label_id": target_label_id,
        "year_id": year_id,
    }
    return data_df, metadata


def _unique_column_values(data_df: pd.DataFrame, column_id: str) -> list[str]:
    if column_id not in data_df.columns:
        return []
    values = (
        data_df[column_id]
        .dropna()
        .astype(str)
        .map(str.strip)
    )
    unique_values = []
    seen = set()
    for value in values:
        if not value or value in seen:
            continue
        seen.add(value)
        unique_values.append(value)
    return unique_values
