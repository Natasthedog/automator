from __future__ import annotations

# Auto-generated split from dash_app.py

import io
import base64
import logging
import copy
from difflib import SequenceMatcher
from dataclasses import dataclass, asdict, is_dataclass
from datetime import date, timedelta
from pathlib import Path
import re
import json
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries, get_column_letter
import numbers
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.parts.embeddedpackage import EmbeddedXlsxPart
from lxml import etree

# NOTE: Reference only. Do NOT import this module in Django.

def download_waterfall_payloads(
    n_clicks,
    data_contents,
    data_name,
    scope_contents,
    scope_name,
    project_name,
    waterfall_targets,
    bucket_data,
):
    if not data_contents or not project_name:
        return no_update, "Please upload the data file and select a project."

    template_path = PROJECT_TEMPLATES.get(project_name)
    if not template_path or not template_path.exists():
        return no_update, "The selected project template could not be found."

    try:
        gathered_df = df_from_contents(data_contents, data_name)
        scope_df = None
        if scope_contents:
            try:
                scope_df = scope_df_from_contents(scope_contents, scope_name)
            except Exception:
                scope_df = None

        prs = Presentation(io.BytesIO(template_path.read_bytes()))
        available_slides = _available_waterfall_template_slides(prs)
        if not available_slides:
            return no_update, "Could not find the <Waterfall Template> slide in the template."
        template_chart = _waterfall_chart_from_slide(
            available_slides[0][1],
            "Waterfall Template",
        )
        if template_chart is None:
            return no_update, "Could not find the waterfall chart on the <Waterfall Template> slide."

        payloads_by_label = compute_waterfall_payloads_for_all_labels(
            gathered_df,
            scope_df,
            bucket_data,
            template_chart,
            target_labels=waterfall_targets,
        )
        payload_json = _waterfall_payloads_to_json(payloads_by_label)
        return dcc.send_string(payload_json, "waterfall_payloads.json"), "Prepared waterfall payload JSON."
    except Exception as exc:
        logger.exception("Waterfall payload JSON generation failed.")
        message = str(exc).strip()
        if not message:
            message = "Unknown error. Check server logs for details."
        return no_update, f"Error ({type(exc).__name__}): {message}"


def generate_deck(
    n_clicks,
    data_contents,
    data_name,
    scope_contents,
    scope_name,
    project_name,
    waterfall_targets,
    bucket_data,
):
    if not data_contents or not project_name:
        return no_update, "Please upload the data file and select a project."

    template_path = PROJECT_TEMPLATES.get(project_name)
    if not template_path or not template_path.exists():
        return no_update, "The selected project template could not be found."
    try:
        df = df_from_contents(data_contents, data_name)
        scope_df = None
        product_description_df = None
        project_details_df = None
        modelled_in_value = None
        metric_value = None
        if scope_contents:
            try:
                scope_df = scope_df_from_contents(scope_contents, scope_name)
            except Exception:
                scope_df = None
            try:
                product_description_df = product_description_df_from_contents(
                    scope_contents, scope_name
                )
            except Exception:
                product_description_df = None
            try:
                project_details_df = project_details_df_from_contents(
                    scope_contents, scope_name
                )
            except Exception:
                project_details_df = None
        if project_details_df is not None:
            modelled_in_value = _project_detail_value_from_df(
                project_details_df,
                "modelled in",
                [
                    "Sales will be modelled in",
                    "Sales will be modeled in",
                    "Sales modelled in",
                    "Sales modeled in",
                ],
                "Sales will be modelled in",
            )
            metric_value = _project_detail_value_from_df(
                project_details_df,
                "metric",
                [
                    "Volume metric (unique per dataset)",
                    "Volume metric unique per dataset",
                    "Volume metric",
                ],
                "Volume metric (unique per dataset)",
            )
        target_brand = target_brand_from_scope_df(scope_df)
        template_bytes = template_path.read_bytes()

        pptx_bytes = build_pptx_from_template(
            template_bytes,
            df,
            target_brand,
            project_name,
            scope_df,
            product_description_df,
            waterfall_targets,
            bucket_data,
            modelled_in_value,
            metric_value,
        )
        return dcc.send_bytes(lambda buff: buff.write(pptx_bytes), "deck.pptx"), "Building deck..."

    except Exception as exc:
        logger.exception("Deck generation failed.")
        message = str(exc).strip()
        if not message:
            message = "Unknown error. Check server logs for details."
        return no_update, f"Error ({type(exc).__name__}): {message}"


def _writer(f):
    pass


def finalize_download(status_text, data_contents):
    # This is a no-op; left for clarity in a larger app. In the minimal example above,
    # you can directly return the 'dcc.send_bytes' with the actual bytes.
    return no_update
